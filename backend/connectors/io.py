"""Format-level readers and primitive normalization for connectors."""

import csv
from decimal import Decimal, InvalidOperation
from pathlib import Path

from backend.connectors.errors import MissingDependencyError, UnsupportedExtensionError


def normalize_header(value) -> str:
    return str(value or "").strip().upper()


def parse_decimal(value) -> Decimal | None:
    if value is None or value == "":
        return None
    if isinstance(value, Decimal):
        return value
    text = str(value).strip()
    negative = text.startswith("(") and text.endswith(")")
    text = (
        text.replace("$", "")
        .replace(",", "")
        .replace("%", "")
        .replace('"', "")
        .replace("(", "")
        .replace(")", "")
        .strip()
    )
    if not text:
        return None
    try:
        number = Decimal(text)
        return -number if negative else number
    except InvalidOperation:
        return None


def read_csv_rows(path: Path):
    with path.open(encoding="utf-8-sig", newline="") as file:
        return list(csv.reader(file))


def read_xlsx_rows(path: Path, *, institution: str):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise MissingDependencyError(
            f"Falta a biblioteca openpyxl para ler o Excel da {institution}."
        ) from exc
    workbook = load_workbook(path, data_only=True, read_only=True)
    worksheet = workbook.active
    return [list(row) for row in worksheet.iter_rows(values_only=True)]


def read_xls_rows(path: Path, *, institution: str):
    try:
        import xlrd
    except ImportError as exc:
        raise MissingDependencyError(
            f"Para ler o arquivo .xls da {institution}, execute uma vez no Terminal: "
            "python3 -m pip install xlrd"
        ) from exc
    workbook = xlrd.open_workbook(path)
    worksheet = workbook.sheet_by_index(0)
    return [worksheet.row_values(index) for index in range(worksheet.nrows)]


def read_tabular_rows(path: Path, *, institution: str, extensions):
    extension = path.suffix.lower()
    if extension not in extensions:
        raise UnsupportedExtensionError(
            f"Formato {institution} não suportado: {extension or '<sem extensão>'}."
        )
    if extension == ".csv":
        return read_csv_rows(path)
    if extension == ".xlsx":
        return read_xlsx_rows(path, institution=institution)
    if extension == ".xls":
        return read_xls_rows(path, institution=institution)
    raise UnsupportedExtensionError(f"Formato não suportado: {extension}")
