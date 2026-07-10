import csv
import re
from pathlib import Path
from typing import Dict, Iterable, List, Sequence

SUPPORTED_EXTENSIONS = {".xls", ".xlsx"}


def _normalize_header(value):
    return str(value or "").strip().upper()


def _parse_number(value):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)

    text = (
        str(value)
        .replace("$", "")
        .replace(",", "")
        .replace("%", "")
        .replace('"', "")
        .strip()
    )
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _csv_rows(path: Path):
    with path.open(encoding="utf-8-sig", newline="") as file:
        return list(csv.reader(file))


def _xlsx_rows(path: Path):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError(
            "Falta a biblioteca openpyxl para ler o Excel do Santander."
        ) from exc

    workbook = load_workbook(path, data_only=True, read_only=True)
    worksheet = workbook.active
    return [list(row) for row in worksheet.iter_rows(values_only=True)]


def _xls_rows(path: Path):
    try:
        import xlrd
    except ImportError as exc:
        raise ValueError(
            "Para ler o arquivo .xls do Santander, execute uma vez no Terminal: python3 -m pip install xlrd"
        ) from exc

    workbook = xlrd.open_workbook(path)
    worksheet = workbook.sheet_by_index(0)
    return [worksheet.row_values(index) for index in range(worksheet.nrows)]


def _read_rows(path: Path):
    extension = path.suffix.lower()
    if extension == ".xlsx":
        return _xlsx_rows(path)
    if extension == ".xls":
        return _xls_rows(path)
    raise ValueError("Formato Santander não suportado. Use XLS ou XLSX.")


def _has_header_markers(row):
    normalized = [_normalize_header(value) for value in row]
    has_value = any(
        "SALDO MOEDA REFERÊNCIA" in value or
        "SALDO NA MOEDA DE REFERÊNCIA" in value
        for value in normalized
    )
    has_weight = any(
        "PESO DA CONTA (%)" in value or "% DO TOTAL" in value
        for value in normalized
    )
    has_name = any(
        value in {
            "NOME DA CARTEIRA",
            "NOME DO ATIVO",
            "ISIN",
            "VALOR DO MERCADO",
            "PREÇO ATUAL",
            "NUMERO DE CONTA",
            "NÚMERO DE CONTA",
            "SALDO",
        }
        for value in normalized
    )
    return has_value and has_weight and has_name


def _is_blank_row(row):
    return all(cell is None or str(cell).strip() == "" for cell in row)


def _is_total_row(row):
    first = str(row[0] or "").strip().upper() if row else ""
    return first.startswith("TOTAL")


def _class_from_block_name(title: str) -> str:
    text = str(title or "").strip().upper()
    if text.startswith("RENDA FIXA"):
        return "Renda Fixa"
    if "AÇÕES" in text:
        return "Ação"
    if "FUNDOS" in text and "RENDA VARIÁVEL" in text:
        return "ETF/Fundo"
    if text.startswith("FUNDOS ALTERNATIVOS"):
        return "Alternativos"
    if text.startswith("FUNDOS DE MERCADOS PRIVADOS"):
        return "Alternativos"
    if text.startswith("LIQUIDEZ"):
        return "Caixa"
    return "Outros"


def _map_block_headers(header):
    mapping = {}
    for index, value in enumerate(header):
        normalized = _normalize_header(value)
        if normalized in {"ISIN", "TICKER", "CÓDIGO", "CODIGO", "TICKER/ISIN"}:
            mapping["symbol"] = index
        elif normalized in {"NOME DA CARTEIRA", "NOME DO ATIVO", "ATIVO", "DESCRIÇÃO", "DESCRICAO"}:
            mapping["description"] = index
        elif normalized in {"SALDO MOEDA REFERÊNCIA", "SALDO NA MOEDA DE REFERÊNCIA", "SALDO"}:
            mapping["value"] = index
        elif normalized in {"PESO DA CONTA (%)", "% DO TOTAL"}:
            mapping["weight"] = index
        elif normalized in {"NUMERO DE CONTA", "NÚMERO DE CONTA", "ACCOUNT", "ACCOUNT NUMBER"}:
            mapping["account"] = index
        elif normalized == "MOEDA":
            mapping.setdefault("currency_candidates", []).append(index)
        elif normalized in {"ASSET CLASS", "CLASS", "CATEGORIA", "CATEGORY"}:
            mapping["asset_class"] = index
        elif normalized in {"VALOR DO MERCADO", "VALOR ESTIMADO", "VALOR INVESTIDO"} and "value" not in mapping:
            mapping["value"] = index
    return mapping


def _safe_get(row, index):
    if index is None or index < 0 or index >= len(row):
        return None
    return row[index]


def _find_currency(row, value_index, currency_indices):
    if currency_indices:
        for index in currency_indices:
            currency = _safe_get(row, index)
            if currency and str(currency).strip().upper() in {"USD", "BRL", "EUR"}:
                return str(currency).strip().upper()
    next_cell = _safe_get(row, value_index + 1)
    if next_cell and str(next_cell).strip().upper() in {"USD", "BRL", "EUR"}:
        return str(next_cell).strip().upper()
    return "USD"


def _looks_like_account_or_category(value) -> bool:
    text = str(value or "").strip().upper()
    if not text or text in {"N/A", "NONE", "NA"}:
        return True
    if "ADVISORY" in text or "ACCOUNT" in text or "CARTEIRA" in text:
        return True
    if text in {"RENDA FIXA", "RENDA VARIÁVEL", "ALTERNATIVOS", "CURTO PRAZO", "LIQUIDEZ", "TOTAL"}:
        return True
    return False


def _extract_coupon(value):
    if value is None:
        return None
    match = re.search(r"(\d+(?:[.,]\d+)?)", str(value))
    if not match:
        return None
    return round(float(match.group(1).replace(",", ".")), 2)


def _extract_year(value):
    if value is None:
        return None
    match = re.search(r"(\d{4})", str(value))
    return match.group(1) if match else None


def _build_display_name(row, mapping, symbol):
    first_cell = str(_safe_get(row, 0) or "").strip()
    fallback_name = str(_safe_get(row, mapping.get("description", 0)) or "").strip()

    if _looks_like_account_or_category(fallback_name):
        fallback_name = first_cell

    if not fallback_name or _looks_like_account_or_category(fallback_name):
        fallback_name = first_cell or symbol or "Ativo"

    coupon = _extract_coupon(_safe_get(row, 2))
    maturity_year = _extract_year(_safe_get(row, 5))

    if coupon is not None and maturity_year:
        issuer = fallback_name or first_cell or symbol or "Ativo"
        return f"{issuer} {coupon:.2f}% {maturity_year}"

    if not fallback_name or _looks_like_account_or_category(fallback_name):
        return symbol or "Ativo"

    return fallback_name


def _find_block_headers(rows):
    headers = []
    for index, row in enumerate(rows):
        if not row:
            continue
        first = str(row[0] or "").strip().upper()
        if _has_header_markers(row) and first:
            headers.append(index)
    return headers


def load_positions(file_path=None):
    path = Path(file_path)
    if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
        raise ValueError("O arquivo Santander deve estar em XLS ou XLSX.")
    if not path.exists():
        raise ValueError(f"Arquivo Santander não encontrado: {path.name}")

    rows = _read_rows(path)
    block_headers = _find_block_headers(rows)
    if not block_headers:
        raise ValueError("Não encontrei blocos Santander com cabeçalhos esperados.")

    positions = []
    for header_index in block_headers:
        header_row = rows[header_index]
        asset_class = _class_from_block_name(header_row[0])
        mapping = _map_block_headers(header_row)
        if "value" not in mapping or "weight" not in mapping:
            continue

        for row in rows[header_index + 1 :]:
            if _is_blank_row(row) or _has_header_markers(row):
                break
            if _is_total_row(row):
                continue

            value = _parse_number(_safe_get(row, mapping["value"]))
            if value is None:
                continue

            symbol = str(_safe_get(row, mapping.get("symbol", -1)) or "").strip()
            if not symbol or symbol.upper() == "N/A":
                symbol = str(_safe_get(row, 0) or "").strip() or ""

            name = _build_display_name(row, mapping, symbol)
            if not name:
                continue

            account = str(_safe_get(row, mapping.get("account", -1)) or "").strip()
            currency = _find_currency(row, mapping["value"], mapping.get("currency_candidates", []))
            weight = _parse_number(_safe_get(row, mapping["weight"]))

            positions.append({
                "institution": "Santander",
                "account": account,
                "symbol": symbol,
                "name": name,
                "description": name,
                "asset_class": asset_class,
                "currency": currency,
                "value": value,
                "weight": weight,
            })

    if not positions:
        raise ValueError("Nenhuma posição Santander foi encontrada no arquivo.")

    return positions
