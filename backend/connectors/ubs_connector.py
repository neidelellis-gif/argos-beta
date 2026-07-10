import csv
from pathlib import Path
from typing import Iterable, List, Sequence

DEFAULT_FILE_PATH = Path(__file__).with_name("UBS_Holdings_08_07_2026.csv")
SUPPORTED_EXTENSIONS = {".csv", ".xls", ".xlsx"}


def classify_asset(symbol, description):
    text = str(description or "").upper()

    if "SAVINGS" in text or "SWEEP" in text:
        return "Caixa"
    if "MATURES" in text or "CALLABLE" in text or "RATE" in text or "NTS" in text:
        return "Renda Fixa"
    if "ETF" in text:
        return "ETF"
    if "FUND" in text:
        return "Fundo"
    if symbol and symbol != "N/A":
        return "Ação"
    return "Outros"


def get_display_name(symbol, description):
    return symbol if symbol and symbol != "N/A" else description


def _normalize_header(value):
    return str(value or "").strip().upper()


def _parse_number(value):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)

    text = (
        str(value)
        .replace("$", "")
        .replace(",", "")
        .replace('"', "")
        .strip()
    )
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _csv_rows(path: Path):
    with path.open(encoding="utf-8-sig", newline="") as file:
        rows = list(csv.reader(file))
    return rows


def _xlsx_rows(path: Path):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError(
            "Falta a biblioteca openpyxl para ler o Excel da UBS."
        ) from exc

    workbook = load_workbook(path, data_only=True, read_only=True)
    worksheet = workbook.active
    return [list(row) for row in worksheet.iter_rows(values_only=True)]


def _xls_rows(path: Path):
    try:
        import xlrd
    except ImportError as exc:
        raise ValueError(
            "Para ler o arquivo .xls da UBS, execute uma vez no Terminal: "
            "python3 -m pip install xlrd"
        ) from exc

    workbook = xlrd.open_workbook(path)
    worksheet = workbook.sheet_by_index(0)
    return [worksheet.row_values(index) for index in range(worksheet.nrows)]


def _read_rows(path: Path):
    extension = path.suffix.lower()
    if extension == ".csv":
        return _csv_rows(path)
    if extension == ".xlsx":
        return _xlsx_rows(path)
    if extension == ".xls":
        return _xls_rows(path)
    raise ValueError(
        "Formato UBS não suportado. Use CSV, XLS ou XLSX."
    )


def _find_header_row(rows):
    required = {"ACCOUNT NUMBER", "DESCRIPTION", "SYMBOL", "VALUE"}

    for index, row in enumerate(rows):
        normalized = {_normalize_header(value) for value in row}
        if required.issubset(normalized):
            return index

    raise ValueError(
        "Não encontrei os cabeçalhos esperados no arquivo UBS."
    )


def load_positions(file_path=None):
    path = Path(file_path) if file_path else DEFAULT_FILE_PATH
    extension = path.suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError("O arquivo UBS deve estar em CSV, XLS ou XLSX.")
    if not path.exists():
        raise ValueError(f"Arquivo UBS não encontrado: {path.name}")

    rows = _read_rows(path)
    header_index = _find_header_row(rows)
    header = [_normalize_header(value) for value in rows[header_index]]

    account_idx = header.index("ACCOUNT NUMBER")
    description_idx = header.index("DESCRIPTION")
    symbol_idx = header.index("SYMBOL")
    value_idx = header.index("VALUE")

    positions = []

    for row in rows[header_index + 1 :]:
        if len(row) <= value_idx:
            continue

        account = str(row[account_idx] or "").strip()
        description = str(row[description_idx] or "").strip()
        symbol = str(row[symbol_idx] or "").strip() or "N/A"
        value = _parse_number(row[value_idx])

        if value is None or not description:
            continue

        positions.append({
            "account": account,
            "symbol": symbol,
            "name": get_display_name(symbol, description),
            "description": description,
            "class": classify_asset(symbol, description),
            "value": value,
        })

    if not positions:
        raise ValueError("Nenhuma posição UBS foi encontrada no arquivo.")

    return positions
